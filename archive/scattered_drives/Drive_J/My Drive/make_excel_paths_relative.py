#!/usr/bin/env python3
"""
Fix the Excel files of a find_duplicates.py output folder so they work from Drive:

  1. Absolute local paths (like F:/data/...) are turned into RELATIVE paths that point to
     each file's location inside the folder (e.g. positive/<dataset>/C1/images/x.jpg).
  2. The "Path in Output" and "Duplicate Of" cells become CLICKABLE HYPERLINKS: clicking
     the path opens the image from the folder you give (for example your Google Drive
     folder  H:/My Drive/De-duplication stages/polypdataset ... withoutduplicates).

Run it on the folder that holds positive_duplicates.xlsx etc. (the Drive copy).
Every .xlsx in that folder is fixed in place. The originals are first copied to a
backup folder NEXT TO it:  <output folder>_excel_backup   (an existing backup is never overwritten)
Running it again is safe (use that to refresh the links if the folder is ever moved).
"""

import os
import re
import shutil
import sys
from urllib.parse import quote

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("This script needs openpyxl.  Install it with:  pip install openpyxl")

CLASS_FOLDERS = {"positive", "negativeonly"}
PATH_HEADER_WORDS = ("path", "duplicate of")    # a header containing one of these is a path column
LINK_HEADER_WORDS = ("path in output", "duplicate of")   # these columns point to files in THIS folder
LINK_FONT = Font(color="0563C1", underline="single")


def looks_absolute(v):
    return bool(re.match(r"^([A-Za-z]:[\\/]|\\\\|/)", v))


def to_relative(value, root):
    """Absolute local path -> relative path with '/'. Already-relative paths are only normalised."""
    if not isinstance(value, str) or not value.strip():
        return value
    norm = value.strip().replace("\\", "/")
    if not looks_absolute(value.strip()):
        return norm
    root_norm = root.replace("\\", "/").rstrip("/")
    if norm.lower().startswith(root_norm.lower() + "/"):          # path starts with this folder
        return norm[len(root_norm) + 1:]
    parts = norm.split("/")
    for i, part in enumerate(parts):                                 # folder was moved: anchor on class folder
        if part.lower() in CLASS_FOLDERS:
            return "/".join(parts[i:])
    return norm                                                      # could not be made relative


def file_uri(folder, rel):
    """file:// link to <folder>/<rel>, percent-encoded the way Excel expects."""
    full = os.path.join(folder, *rel.split("/")).replace("\\", "/")
    prefix = "file://" if full.startswith("/") else "file:///"
    return prefix + quote(full, safe="/:()")


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else input("Enter the OUTPUT folder (with the Excel files): ")
    folder = os.path.abspath(os.path.expanduser(folder.strip().strip('"')))
    if not os.path.isdir(folder):
        sys.exit(f"Error: '{folder}' is not a valid folder.")

    xlsx = sorted(f for f in os.listdir(folder) if f.lower().endswith(".xlsx") and not f.startswith("~$"))
    if not xlsx:
        sys.exit("No .xlsx files found in that folder.")

    backup = folder.rstrip("\\/") + "_excel_backup"
    os.makedirs(backup, exist_ok=True)

    for name in xlsx:
        path = os.path.join(folder, name)
        if not os.path.exists(os.path.join(backup, name)):           # keep the very first original
            shutil.copy2(path, os.path.join(backup, name))

        wb = load_workbook(path)
        ws = wb.active
        headers = [str(c.value or "").lower() for c in ws[1]]
        cols = [i for i, h in enumerate(headers) if any(w in h for w in PATH_HEADER_WORDS)]
        if not cols:
            print(f"{name}: no path columns, skipped")
            continue

        link_cols = {i for i in cols if any(w in headers[i] for w in LINK_HEADER_WORDS)}
        changed = unresolved = links = 0
        for row in ws.iter_rows(min_row=2):
            for i in cols:
                cell = row[i]
                new = to_relative(cell.value, folder)
                if new != cell.value:
                    cell.value = new
                    changed += 1
                if isinstance(new, str) and looks_absolute(new):
                    unresolved += 1
                elif i in link_cols and isinstance(new, str) and new.strip():
                    cell.hyperlink = file_uri(folder, new)
                    cell.font = LINK_FONT
                    links += 1

        for i in range(1, ws.max_column + 1):                         # re-fit column widths
            letter = get_column_letter(i)
            width = max(len(str(c.value)) if c.value is not None else 0 for c in ws[letter])
            ws.column_dimensions[letter].width = min(width + 2, 90)
        wb.save(path)
        msg = f"{name}: {changed} paths made relative, {links} hyperlinks added"
        if unresolved:
            msg += f", {unresolved} could not be converted"
        print(msg)
        if links > 60000:
            print("  Note: Excel handles at most about 65,000 hyperlinks per sheet.")

    print(f"\nDone. Original sheets backed up in: {backup}")


if __name__ == "__main__":
    main()
