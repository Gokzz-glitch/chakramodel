import collections
import collections.abc
from pptx import Presentation

def replace_text_in_shape(shape, old_text, new_text):
    if not hasattr(shape, "text_frame"):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text)

def set_shape_text(shape, new_text):
    if not hasattr(shape, "text_frame"):
        return
    # This preserves the paragraph styling of the first paragraph/run
    if len(shape.text_frame.paragraphs) > 0:
        p = shape.text_frame.paragraphs[0]
        # Clear all runs
        for run in p.runs:
            run.text = ""
        # Add new text to the first run if exists, else create one
        if len(p.runs) > 0:
            p.runs[0].text = new_text
        else:
            run = p.add_run()
            run.text = new_text
        # Remove extra paragraphs
        for i in range(len(shape.text_frame.paragraphs)-1, 0, -1):
            p = shape.text_frame.paragraphs[i]
            p._element.getparent().remove(p._element)
    else:
        shape.text = new_text

def update_presentation(template_path, output_path):
    prs = Presentation(template_path)

    # Global Replacements
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                replace_text_in_shape(shape, "[Project Name]", "ChakraModel")
                replace_text_in_shape(shape, "[Hackathon Name] · [Date]", "Final Hackathon Pitch Deck")

    # Slide 1: Cover
    s1 = prs.slides[0]
    set_shape_text(s1.shapes[37], "Edge-AI Polyp Detection with Conformal Uncertainty Quantification")
    set_shape_text(s1.shapes[38], "ChakraModel")
    set_shape_text(s1.shapes[39], "Edge-AI Polyp Detection with Conformal Uncertainty Quantification")

    # Slide 2: Team
    s2 = prs.slides[1]
    set_shape_text(s2.shapes[43], "GOKUL R — AI/ML Engineer")
    set_shape_text(s2.shapes[46], "Varsha Rao — BME")
    set_shape_text(s2.shapes[49], "Jayasree — BME")
    set_shape_text(s2.shapes[52], " ")
    set_shape_text(s2.shapes[53], "Mentor / Guide: Gladys — AI/ML Professor")

    # Slide 3: Problem
    s3 = prs.slides[2]
    set_shape_text(s3.shapes[42], "CRC is the 2nd leading global cause of cancer deaths (4th in India with 83k cases).")
    set_shape_text(s3.shapes[44], "Endoscopists (miss 20-25% of adenomas).")
    set_shape_text(s3.shapes[47], "Every 1% drop in ADR increases CRC risk by 3%.")
    set_shape_text(s3.shapes[50], "High miss rates in standard colonoscopies due to flicker and human error.")

    # Slide 4: Solution
    s4 = prs.slides[3]
    set_shape_text(s4.shapes[42], "[ 5-Stage Architecture Diagram Here ]")
    set_shape_text(s4.shapes[43], "5-stage AI pipeline (YOLOv11m + BoT-SORT + Split-Conformal Prediction) for real-time Edge-AI Polyp Detection.")
    set_shape_text(s4.shapes[44], "Both Edge and SaaS deployment models.")

    # Slide 5: Edge (Why it's better)
    s5 = prs.slides[4]
    set_shape_text(s5.shapes[43], "Zero Alert Flicker (≤2%)")
    set_shape_text(s5.shapes[44], "Stable bounding boxes and BoT-SORT tracking for endoscopists.")
    set_shape_text(s5.shapes[47], "100% Offline Edge Privacy")
    set_shape_text(s5.shapes[48], "No patient data leaves the hospital network.")
    set_shape_text(s5.shapes[51], "FDA-Aligned UQ")
    set_shape_text(s5.shapes[52], "Split-Conformal prediction provides a mathematical 95% coverage guarantee.")
    set_shape_text(s5.shapes[55], "Hybrid Edge & SaaS")
    set_shape_text(s5.shapes[56], "Flexible deployments: Jetson AGX Orin for Edge, API for Cloud SaaS.")

    # Slide 6: Tech Stack
    s6 = prs.slides[5]
    set_shape_text(s6.shapes[42], "Hardware")
    set_shape_text(s6.shapes[43], "Jetson AGX Orin")
    set_shape_text(s6.shapes[44], "Frameworks")
    set_shape_text(s6.shapes[45], "PyTorch, TensorRT")
    set_shape_text(s6.shapes[46], "Models")
    set_shape_text(s6.shapes[47], "YOLOv11m, BoT-SORT")
    set_shape_text(s6.shapes[48], "Validation")
    set_shape_text(s6.shapes[49], "MAPIE (Split-Conformal Prediction)")
    set_shape_text(s6.shapes[50], "")
    set_shape_text(s6.shapes[51], "")
    set_shape_text(s6.shapes[52], "")
    set_shape_text(s6.shapes[53], "Why these choices: Enables robust 5-stage medical AI pipeline in 24 hours.")

    # Slide 7: Roadmap
    s7 = prs.slides[6]
    set_shape_text(s7.shapes[44], "Offline pipeline with Edge and SaaS foundation.")
    set_shape_text(s7.shapes[47], "Live HDMI testing in clinical setup.")
    set_shape_text(s7.shapes[50], "Clinical trials and enterprise SaaS launch.")

    # Slide 8: Demo
    s8 = prs.slides[7]
    set_shape_text(s8.shapes[42], "[ Demo Image: YOLO bounding box with 95% UQ score ]")
    set_shape_text(s8.shapes[44], "[ Demo Image: BoT-SORT tracking consistency across frames ]")
    set_shape_text(s8.shapes[46], "[ Demo Image: SaaS API & Edge Dashboard ]")
    set_shape_text(s8.shapes[47], "Live demo: running inference locally on Jetson / SaaS Cloud.")

    # Slide 9: Impact
    s9 = prs.slides[8]
    set_shape_text(s9.shapes[43], "Saves lives by increasing Adenoma Detection Rate (ADR).")
    set_shape_text(s9.shapes[45], "Endoscopy clinics globally (B2B).")
    set_shape_text(s9.shapes[47], "Reduces 3% CRC risk for every 1% ADR increase.")
    set_shape_text(s9.shapes[50], "Hospital networks and telemedicine platforms.")
    set_shape_text(s9.shapes[52], "Eliminates operational cloud costs via Edge processing.")
    set_shape_text(s9.shapes[54], "Dual SaaS and Edge licensing models for rapid scaling.")

    # Slide 10: USP (Scientific Criteria) - Highly Emphasized
    s10 = prs.slides[9]
    # Highlight Split-Conformal Prediction giving a mathematical 95% coverage guarantee, directly aligning with FDA CDRH 2024 AI guidelines.
    set_shape_text(s10.shapes[43], "Split-Conformal Prediction")
    set_shape_text(s10.shapes[44], "Provides a mathematical 95% coverage guarantee.")
    set_shape_text(s10.shapes[47], "FDA CDRH 2024 AI Guidelines")
    set_shape_text(s10.shapes[48], "Directly aligns with strict clinical UQ requirements.")
    set_shape_text(s10.shapes[51], "Zero Alert Flicker")
    set_shape_text(s10.shapes[52], "Advanced BoT-SORT integration for stable UI.")
    set_shape_text(s10.shapes[55], "Edge-First Privacy")
    set_shape_text(s10.shapes[56], "Real-time inference without cloud dependency.")

    # Slide 11: Marketing (Business)
    s11 = prs.slides[10]
    set_shape_text(s11.shapes[43], "B2B Hardware + SaaS Licensing")
    set_shape_text(s11.shapes[45], "Medical conferences & gastroenterology partnerships")
    set_shape_text(s11.shapes[47], "Direct sales to clinics, API access for telemedicine platforms")
    set_shape_text(s11.shapes[50], "Endoscopy clinics & large hospital networks")
    set_shape_text(s11.shapes[52], "FDA-aligned Edge-AI Polyp Detection")
    set_shape_text(s11.shapes[54], "High reliability and zero-flicker UI drives retention")

    # Slide 12: Revenue Growth
    s12 = prs.slides[11]
    set_shape_text(s12.shapes[43], "B2B Hardware Sales (Jetson AGX Orin bundles)")
    set_shape_text(s12.shapes[45], "SaaS Cloud API Subscriptions")
    set_shape_text(s12.shapes[47], "Tiered licensing: Edge (CapEx + Maintenance) vs SaaS (OpEx)")
    set_shape_text(s12.shapes[50], "Penetrate 5% of early-adopter endoscopy clinics")
    set_shape_text(s12.shapes[52], "Leverage the 25% CAGR AI Endoscopy Market")
    set_shape_text(s12.shapes[54], "Low operational cloud costs due to Edge-first inference")

    # Slide 13: Sustainability
    s13 = prs.slides[12]
    set_shape_text(s13.shapes[43], "Small ML Ops team for model updates")
    set_shape_text(s13.shapes[45], "Low maintenance (Inference strictly on the Edge)")
    set_shape_text(s13.shapes[47], "Improved patient outcomes with minimal carbon footprint")
    set_shape_text(s13.shapes[50], "B2B sales and potential medical tech grants")
    set_shape_text(s13.shapes[52], "Very low ongoing cost vs high licensing revenue")
    set_shape_text(s13.shapes[54], "Break-even at 15 clinic deployments")

    # Slide 14: Budget - Highly Emphasized
    s14 = prs.slides[13]
    # Total Prototype BOM is ~$2,500 ($1,999 Jetson AGX Orin, $150 Capture card, $351 Cloud training)
    set_shape_text(s14.shapes[43], "Jetson AGX Orin Edge Hardware — $1,999")
    set_shape_text(s14.shapes[45], "Cloud Training Compute — $351")
    set_shape_text(s14.shapes[47], "HDMI Capture Card & Materials — $150")
    set_shape_text(s14.shapes[50], "Self-funded Prototype Build")
    set_shape_text(s14.shapes[52], "Total Prototype BOM — ~$2,500")
    set_shape_text(s14.shapes[54], "Remaining / contingency — $0")

    # Slide 15: Achievements
    s15 = prs.slides[14]
    set_shape_text(s15.shapes[43], "Rapid Development")
    set_shape_text(s15.shapes[44], "Built a 5-stage medical AI pipeline in 24 hours.")
    set_shape_text(s15.shapes[47], "Advanced UQ")
    set_shape_text(s15.shapes[48], "Implemented Split-Conformal Prediction.")
    set_shape_text(s15.shapes[51], "Zero-Flicker Tracking")
    set_shape_text(s15.shapes[52], "Successfully integrated BoT-SORT for stable UI.")
    set_shape_text(s15.shapes[55], "Deployment Ready")
    set_shape_text(s15.shapes[56], "Configured for both Edge and SaaS.")

    # Slide 16: Thank You
    s16 = prs.slides[15]
    set_shape_text(s16.shapes[38], "ChakraModel")
    set_shape_text(s16.shapes[42], "Gokul R, Varsha Rao, Jayasree")
    set_shape_text(s16.shapes[44], "ChakraModelTeam@example.com")
    set_shape_text(s16.shapes[46], "github.com/ChakraModel")

    prs.save(output_path)
    print("Successfully created ChakraModel_Full_Pitch.pptx")

if __name__ == "__main__":
    update_presentation(r"C:\Users\imgk3\Downloads\hackathon_pitch_deck_template.pptx", r"C:\Users\imgk3\Downloads\ChakraModel_Full_Pitch.pptx")
