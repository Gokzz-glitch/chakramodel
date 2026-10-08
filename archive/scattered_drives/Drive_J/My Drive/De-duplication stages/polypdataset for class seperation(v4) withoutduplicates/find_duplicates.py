#!/usr/bin/env python3
"""
Duplicate finder for the RENAMED dataset folder.

Asks for:
  1. the renamed folder (contains `positive` and `negativeonly`)
  2. a NEW, empty output folder

Duplicates are detected bit-for-bit (SHA-256, then a byte-by-byte confirmation)
separately inside `positive` and inside `negativeonly`.
For every group of identical images the first one (natural path order) is kept
as the original; the others are the duplicates.

Output folder = same structure as the renamed folder, plus:
  positive/duplicates of positive/          -> duplicate images AND their masks
  negativeonly/duplicates of negative/      -> duplicate images
(each with the same sub-folder directory as before)

Excel files written into the output folder:
  positive_duplicates.xlsx      positive_originals.xlsx
  negativeonly_duplicates.xlsx  negativeonly_originals.xlsx
  (+ cross_class_duplicates.xlsx only if an identical image exists in BOTH classes)

The renamed folder itself is never modified (set COPY_FILES = False to move instead).
"""

import filecmp
import hashlib
import os
import re
import shutil
import sys
from collections import defaultdict

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("This script needs openpyxl.  Install it with:  pip install openpyxl")

# ----------------------------- settings -----------------------------
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".gif", ".webp"}

# class folder (lowercase)  ->  name of the extra duplicates folder inside it
CLASSES = {
    "positive": "duplicates of positive",
    "negativeonly": "duplicates of negative",
}

# Masks may be named like their image plus one of these endings (e.g. img1.jpg -> img1_mask.jpg)
MASK_SUFFIXES = ("_mask", "-mask")

SKIP_ROOT_FILES = {"file_directory.xlsx"}   # old directory sheet (paths would be stale)
COPY_FILES = True                           # True = copy (safe), False = move
# --------------------------------------------------------------------


def natural_key(text):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", text.lower())]


def is_image(name):
    return os.path.splitext(name)[1].lower() in IMAGE_EXTS


def in_masks(rel):
    return len(rel) >= 2 and rel[-2].lower() in ("masks", "mask")


def class_of(rel):
    """'positive' / 'negativeonly' / None for a file's relative path parts."""
    return rel[0].lower() if len(rel) > 1 and rel[0].lower() in CLASSES else None


def sha256_of(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def hash_groups(paths):
    """Groups (lists) of files whose full contents hash the same. Only files that
    share a size are hashed, which skips most of the reading."""
    by_size = defaultdict(list)
    for p in paths:
        by_size[os.path.getsize(p)].append(p)
    candidates = [p for lst in by_size.values() if len(lst) > 1 for p in lst]
    print(f"Hashing {len(candidates)} files that share a size with another file...")
    groups, done = [], 0
    for lst in by_size.values():
        if len(lst) < 2:
            continue
        by_hash = defaultdict(list)
        for p in lst:
            by_hash[sha256_of(p)].append(p)
            done += 1
            if done % 2000 == 0:
                print(f"  ...{done} hashed")
        groups += [g for g in by_hash.values() if len(g) > 1]
    return groups


def save_excel(path, headers, rows, title):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(["No."] + headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, r in enumerate(rows, start=1):
        ws.append([i, *r])
    for col in range(1, len(headers) + 2):
        letter = get_column_letter(col)
        width = max(len(str(c.value)) if c.value is not None else 0 for c in ws[letter])
        ws.column_dimensions[letter].width = min(width + 2, 90)
    ws.freeze_panes = "A2"
    wb.save(path)


def main():
    if len(sys.argv) >= 3:
        src, dst = sys.argv[1], sys.argv[2]
    else:
        src = input("Enter the RENAMED folder path: ")
        dst = input("Enter the NEW output folder path: ")
    src = os.path.abspath(os.path.expanduser(src.strip().strip('"')))
    dst = os.path.abspath(os.path.expanduser(dst.strip().strip('"')))

    if not os.path.isdir(src):
        sys.exit(f"Error: '{src}' is not a valid folder.")
    try:
        if os.path.commonpath([src, dst]) == src:
            sys.exit("Error: the output folder must be outside the renamed folder.")
    except ValueError:
        pass
    if os.path.isdir(dst) and os.listdir(dst):
        sys.exit("Error: the output folder is not empty. Give a new or empty folder.")

    # ---------------- 1. scan ----------------
    files, dirs = [], []                      # files: (abs path, rel parts)
    for dirpath, dirnames, filenames in os.walk(src):
        dirnames.sort(key=natural_key)
        rel = os.path.relpath(dirpath, src)
        parts = [] if rel == "." else rel.split(os.sep)
        dirs.append(parts)
        for f in sorted(filenames, key=natural_key):
            if not parts and f in SKIP_ROOT_FILES:
                continue
            files.append((os.path.join(dirpath, f), parts + [f]))
    rel_of = dict(files)

    top_dirs = {p[0].lower(): p[0] for p in dirs if len(p) == 1}
    if not any(k in top_dirs for k in CLASSES):
        sys.exit("Error: could not find a 'positive' or 'negativeonly' folder inside the renamed folder.")

    to_check = [p for p, rel in files
                if class_of(rel) and is_image(rel[-1]) and not in_masks(rel)]
    print(f"Found {len(files)} files, {len(to_check)} images to check for duplicates.")

    # ---------------- 2. find duplicates (bit level) ----------------
    dup_of, cross = {}, []                    # duplicate -> kept original
    for g in hash_groups(to_check):
        by_cls = defaultdict(list)
        for p in g:
            by_cls[class_of(rel_of[p])].append(p)
        firsts = []
        for lst in by_cls.values():
            lst.sort(key=lambda p: natural_key("/".join(rel_of[p])))
            firsts.append(lst[0])
            for d in lst[1:]:
                if filecmp.cmp(lst[0], d, shallow=False):     # byte-by-byte confirmation
                    dup_of[d] = lst[0]
        if len(by_cls) > 1 and filecmp.cmp(firsts[0], firsts[1], shallow=False):
            cross.append(g)

    # ---------------- 3. pair duplicate images with their masks ----------------
    mask_cache = {}

    def find_mask(img_path):
        parent = os.path.dirname(img_path)
        if os.path.basename(parent).lower() != "images":
            return None
        gp = os.path.dirname(parent)
        if gp not in mask_cache:
            idx = {"name": {}, "stem": {}, "stripped": {}}
            for d in os.listdir(gp):
                full = os.path.join(gp, d)
                if d.lower() in ("masks", "mask") and os.path.isdir(full):
                    for f in os.listdir(full):
                        idx["name"][f] = os.path.join(full, f)
                        s = os.path.splitext(f)[0]
                        idx["stem"].setdefault(s, os.path.join(full, f))
                        for suf in MASK_SUFFIXES:
                            if s.endswith(suf):
                                idx["stripped"].setdefault(s[:-len(suf)], os.path.join(full, f))
            mask_cache[gp] = idx
        name = os.path.basename(img_path)
        idx = mask_cache[gp]
        base = os.path.splitext(name)[0]
        return idx["name"].get(name) or idx["stem"].get(base) or idx["stripped"].get(base)

    dup_mask_of = {}                          # duplicate mask -> its duplicate image
    keep_mask_of = {}                         # duplicate mask -> mask of the kept original
    no_mask = defaultdict(int)
    for d, keep in dup_of.items():
        m = find_mask(d)
        if m:
            dup_mask_of[m] = d
            keep_mask_of[m] = find_mask(keep)
        elif class_of(rel_of[d]) == "positive":
            no_mask[d] = 1

    # ---------------- 4. copy / move into the output folder ----------------
    def dest_for(rel, is_dup):
        if is_dup:
            return os.path.join(dst, rel[0], CLASSES[rel[0].lower()], *rel[1:])
        return os.path.join(dst, *rel)

    os.makedirs(dst, exist_ok=True)
    for parts in dirs:                        # same folder structure as the renamed folder
        os.makedirs(os.path.join(dst, *parts), exist_ok=True)
    for key, name in CLASSES.items():         # the extra duplicates folders
        if key in top_dirs:
            os.makedirs(os.path.join(dst, top_dirs[key], name), exist_ok=True)

    transfer = shutil.copy2 if COPY_FILES else shutil.move
    rows = {k: {"dup": [], "orig": []} for k in CLASSES}
    for n, (p, rel) in enumerate(files, start=1):
        key = class_of(rel)
        is_dup = p in dup_of or p in dup_mask_of
        dest = dest_for(rel, is_dup)
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        transfer(p, dest)
        if key and is_image(rel[-1]):
            kind = "mask" if in_masks(rel) else "image"
            if is_dup:
                kept = dup_of.get(p) or keep_mask_of.get(p)
                kept_dest = dest_for(rel_of[kept], False) if kept in rel_of else ""
                rows[key]["dup"].append((kind, rel[-1], dest, p, kept_dest))
            else:
                rows[key]["orig"].append((kind, rel[-1], dest, p))
        if n % 2000 == 0:
            print(f"  ...{n} files transferred")

    # ---------------- 5. Excel sheets ----------------
    dup_headers = ["Type", "File Name", "Path in Output", "Path in Renamed Folder",
                   "Duplicate Of (kept original, in output)"]
    org_headers = ["Type", "File Name", "Path in Output", "Path in Renamed Folder"]
    for key in CLASSES:
        save_excel(os.path.join(dst, f"{key}_duplicates.xlsx"), dup_headers, rows[key]["dup"], "Duplicates")
        save_excel(os.path.join(dst, f"{key}_originals.xlsx"), org_headers, rows[key]["orig"], "Originals")
    if cross:
        cross_rows = [(i, class_of(rel_of[p]), p) for i, g in enumerate(cross, 1)
                      for p in sorted(g, key=lambda x: natural_key("/".join(rel_of[x])))]
        save_excel(os.path.join(dst, "cross_class_duplicates.xlsx"),
                   ["Group", "Class", "Path in Renamed Folder"], cross_rows, "Cross-class")

    # ---------------- 6. verify against the real output folder ----------------
    def count_images(root):
        return sum(1 for _, _, fs in os.walk(root) for f in fs if is_image(f))

    print("\n" + "=" * 62)
    print("RESULT")
    print("=" * 62)
    all_ok = True
    total_moved = 0
    for key, label in (("positive", "POSITIVE"), ("negativeonly", "NEGATIVEONLY")):
        if key not in top_dirs:
            continue
        d_imgs = sum(1 for r in rows[key]["dup"] if r[0] == "image")
        d_masks = sum(1 for r in rows[key]["dup"] if r[0] == "mask")
        o_imgs = sum(1 for r in rows[key]["orig"] if r[0] == "image")
        o_masks = sum(1 for r in rows[key]["orig"] if r[0] == "mask")
        src_total = sum(1 for p, rel in files if class_of(rel) == key and is_image(rel[-1]))
        out_total = count_images(os.path.join(dst, top_dirs[key]))
        out_dups = count_images(os.path.join(dst, top_dirs[key], CLASSES[key]))
        ok = (src_total == out_total == d_imgs + d_masks + o_imgs + o_masks) and out_dups == d_imgs + d_masks
        all_ok &= ok
        total_moved += d_imgs + d_masks
        print(f"\n{label}")
        print(f"  Images checked for duplicates : {o_imgs + d_imgs}")
        print(f"  Duplicate images moved        : {d_imgs}")
        if key == "positive":
            print(f"  Duplicate masks moved         : {d_masks}")
        print(f"  Original images kept          : {o_imgs}")
        if key == "positive":
            print(f"  Original masks kept           : {o_masks}")
        print(f"  Total moved into duplicates   : {d_imgs + d_masks}")
        print(f"  Check (input = output)        : {src_total} -> {out_total}  {'OK' if ok else 'MISMATCH!'}")

    print(f"\nTOTAL MOVED (positive + negativeonly): {total_moved}")
    if no_mask:
        print(f"WARNING: {len(no_mask)} duplicate positive image(s) had no matching mask.")
    if cross:
        print(f"WARNING: {len(cross)} image(s) are identical across positive AND negativeonly "
              f"(left in place, listed in cross_class_duplicates.xlsx).")
    print(f"\nOutput folder: {dst}")
    print("Excel files : positive_duplicates / positive_originals / "
          "negativeonly_duplicates / negativeonly_originals")
    if not all_ok:
        print("\n*** Counts do not add up - please check the output folder. ***")


if __name__ == "__main__":
    main()
