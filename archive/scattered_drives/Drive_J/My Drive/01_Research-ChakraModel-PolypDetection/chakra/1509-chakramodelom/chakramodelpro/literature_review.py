"""
================================================================================
 LITERATURE REVIEW AUTOMATION SCRIPT
================================================================================
Ithu enna pannum:
  1. Keyword kudutha, Semantic Scholar + arXiv la papers search pannum
     (Google Scholar scrape panna mudiyathu - blocked/against ToS, so idhu
      reliable + official alternative, and citation count (impact score) um
      free ah kudukum)
  2. Gemini API use panni, retrieved papers ah subtopics ah cluster pannum
     (e.g. "earthquake" -> 1. Bird movement, 2. Post-tsunami poverty, etc.)
  3. Ovvoru paper ku abstract padichi Gemini extract pannum:
        - Dataset used
        - Model / Architecture
        - Layers / key components
        - Base paper it builds on
        - Whether it's actually ML-related (true/false)
  4. Ellame Excel file ah export pannum - clickable links + impact score sort

SETUP (run once):
  pip install requests openpyxl google-generativeai --break-system-packages
  export GEMINI_API_KEY="your-gemini-api-key-here"

USAGE:
  python literature_review.py "earthquake"
  python literature_review.py "earthquake" --limit 40 --ml-only
  python literature_review.py "colonoscopy polyp segmentation" --limit 25

OUTPUT:
  A file like literature_review_earthquake.xlsx in the current folder.
================================================================================
"""

import os
import sys
import re
import json
import time
import argparse
import requests
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

try:
    import google.generativeai as genai
except ImportError:
    genai = None


SEMANTIC_SCHOLAR_API = "https://api.semanticscholar.org/graph/v1/paper/search"
ARXIV_API = "http://export.arxiv.org/api/query"


# ---------------------------------------------------------------------------
# STEP 1: FETCH PAPERS
# ---------------------------------------------------------------------------

def fetch_semantic_scholar(keyword, limit=50):
    """Free API, no key needed. Gives citation counts -> used as impact score."""
    params = {
        "query": keyword,
        "limit": min(limit, 100),
        "fields": "title,abstract,year,citationCount,authors,externalIds,url,venue",
    }
    papers = []
    try:
        resp = requests.get(SEMANTIC_SCHOLAR_API, params=params, timeout=25)
        resp.raise_for_status()
        for p in resp.json().get("data", []):
            if not p.get("abstract"):
                continue
            papers.append({
                "title": p.get("title", "Untitled").strip(),
                "abstract": p.get("abstract", "").strip(),
                "year": p.get("year"),
                "citations": p.get("citationCount", 0) or 0,
                "authors": ", ".join(a["name"] for a in p.get("authors", [])[:3]),
                "url": p.get("url") or "",
                "venue": p.get("venue", ""),
                "source": "Semantic Scholar",
            })
    except requests.RequestException as e:
        print(f"[warn] Semantic Scholar fetch failed: {e}")
    return papers


def fetch_arxiv(keyword, limit=30):
    """Good for recent ML/CS preprints. No citation count available here."""
    params = {
        "search_query": f'all:"{keyword}"',
        "start": 0,
        "max_results": limit,
    }
    papers = []
    try:
        resp = requests.get(ARXIV_API, params=params, timeout=25)
        resp.raise_for_status()
        ns = {"atom": "http://www.w3.org/2005/Atom"}
        root = ET.fromstring(resp.text)
        for entry in root.findall("atom:entry", ns):
            title = entry.find("atom:title", ns).text.strip().replace("\n", " ")
            summary = entry.find("atom:summary", ns).text.strip().replace("\n", " ")
            link = entry.find("atom:id", ns).text.strip()
            published = entry.find("atom:published", ns).text[:4]
            authors = ", ".join(
                a.find("atom:name", ns).text for a in entry.findall("atom:author", ns)[:3]
            )
            papers.append({
                "title": title,
                "abstract": summary,
                "year": int(published) if published.isdigit() else None,
                "citations": None,
                "authors": authors,
                "url": link,
                "venue": "arXiv",
                "source": "arXiv",
            })
    except Exception as e:
        print(f"[warn] arXiv fetch failed: {e}")
    return papers


def dedupe(papers):
    seen, out = set(), []
    for p in papers:
        key = re.sub(r"[^a-z0-9]", "", p["title"].lower())[:60]
        if key and key not in seen:
            seen.add(key)
            out.append(p)
    return out


# ---------------------------------------------------------------------------
# STEP 2: GEMINI HELPERS
# ---------------------------------------------------------------------------

def get_gemini_model(api_key):
    if genai is None:
        sys.exit("google-generativeai not installed. Run: pip install google-generativeai --break-system-packages")
    genai.configure(api_key=api_key)
    return genai.GenerativeModel("gemini-2.0-flash")


def safe_json_extract(text):
    """Gemini sometimes wraps JSON in ```json fences - strip them before parsing."""
    cleaned = re.sub(r"^```(json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        return None


def cluster_into_subtopics(model, topic, papers):
    """Ask Gemini to group papers into numbered subtopics."""
    listing = "\n".join(f"{i}: {p['title']} -- {p['abstract'][:200]}" for i, p in enumerate(papers))
    prompt = f"""
You are helping organize a literature review on the topic: "{topic}".
Below is a numbered list of papers (index: title -- short abstract).

{listing}

Group these papers into meaningful research subtopics under "{topic}".
Respond with ONLY valid JSON, no markdown fences, in this exact shape:
{{
  "subtopics": [
    {{"id": 1, "name": "Short subtopic name", "paper_indices": [0, 3, 7]}},
    {{"id": 2, "name": "Another subtopic", "paper_indices": [1, 2]}}
  ]
}}
Every paper index must appear in exactly one subtopic. Keep subtopic names short (3-6 words).
"""
    resp = model.generate_content(prompt)
    data = safe_json_extract(resp.text)
    if not data:
        # fallback: everything in one bucket if Gemini output couldn't be parsed
        return [{"id": 1, "name": topic, "paper_indices": list(range(len(papers)))}]
    return data["subtopics"]


def extract_ml_details(model, paper):
    """Ask Gemini to pull structured ML info out of a single abstract."""
    prompt = f"""
Read this paper's title and abstract. Extract ML-specific details if present.

Title: {paper['title']}
Abstract: {paper['abstract']}

Respond with ONLY valid JSON, no markdown fences, in this exact shape:
{{
  "is_ml_related": true/false,
  "dataset": "dataset name(s) used, or 'Not mentioned'",
  "architecture": "model/architecture name, or 'Not mentioned'",
  "layers": "key layers/components mentioned, or 'Not mentioned'",
  "base_paper": "prior work this builds on, or 'Not mentioned'"
}}
"""
    try:
        resp = model.generate_content(prompt)
        data = safe_json_extract(resp.text)
        if data:
            return data
    except Exception as e:
        print(f"[warn] Gemini extraction failed for '{paper['title'][:40]}...': {e}")
    return {
        "is_ml_related": False,
        "dataset": "Not extracted",
        "architecture": "Not extracted",
        "layers": "Not extracted",
        "base_paper": "Not extracted",
    }


# ---------------------------------------------------------------------------
# STEP 3: EXCEL EXPORT
# ---------------------------------------------------------------------------

def export_to_excel(topic, subtopics, papers, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Literature Review"

    headers = [
        "Subtopic #", "Subtopic", "Title", "Year", "Authors",
        "Impact Score (Citations)", "ML-Related", "Dataset",
        "Architecture", "Layers", "Base Paper", "Source", "Link",
    ]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    rows = []
    for sub in subtopics:
        for idx in sub["paper_indices"]:
            p = papers[idx]
            rows.append((sub, p))

    # sort by subtopic id, then by impact score descending
    rows.sort(key=lambda r: (r[0]["id"], -(r[1].get("citations") or 0)))

    for sub, p in rows:
        row = [
            sub["id"],
            sub["name"],
            p["title"],
            p.get("year") or "",
            p.get("authors") or "",
            p.get("citations") if p.get("citations") is not None else "N/A",
            "Yes" if p.get("is_ml_related") else "No",
            p.get("dataset", ""),
            p.get("architecture", ""),
            p.get("layers", ""),
            p.get("base_paper", ""),
            p.get("source", ""),
            p.get("url", ""),
        ]
        ws.append(row)
        r = ws.max_row
        link_cell = ws.cell(row=r, column=13)
        if p.get("url"):
            link_cell.hyperlink = p["url"]
            link_cell.value = "Open paper"
            link_cell.font = Font(color="0563C1", underline="single")
        # highlight high-impact papers
        citations = p.get("citations") or 0
        if citations >= 50:
            ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor="C6EFCE")
        elif citations >= 10:
            ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor="FFEB9C")

    widths = [10, 22, 40, 6, 22, 14, 10, 25, 25, 25, 25, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    wb.save(filename)
    print(f"\n[done] Saved: {filename}")


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Automated literature review builder")
    parser.add_argument("topic", help="Research keyword/topic, e.g. 'earthquake'")
    parser.add_argument("--limit", type=int, default=30, help="Max papers per source")
    parser.add_argument("--ml-only", action="store_true", help="Keep only ML-related papers")
    parser.add_argument("--api-key", default=os.environ.get("GEMINI_API_KEY"),
                         help="Gemini API key (or set GEMINI_API_KEY env var)")
    args = parser.parse_args()

    if not args.api_key:
        sys.exit("Gemini API key missing. Pass --api-key or set GEMINI_API_KEY env var.")

    print(f"[1/4] Fetching papers on '{args.topic}' ...")
    papers = fetch_semantic_scholar(args.topic, args.limit) + fetch_arxiv(args.topic, args.limit)
    papers = dedupe(papers)
    print(f"      -> {len(papers)} unique papers with abstracts found.")

    if not papers:
        sys.exit("No papers found. Try a broader keyword.")

    model = get_gemini_model(args.api_key)

    print("[2/4] Extracting ML details from each abstract via Gemini ...")
    for i, p in enumerate(papers, 1):
        details = extract_ml_details(model, p)
        p.update(details)
        print(f"      ({i}/{len(papers)}) {p['title'][:60]}...")
        time.sleep(1.2)  # gentle pacing to avoid rate limits on free tier

    if args.ml_only:
        before = len(papers)
        papers = [p for p in papers if p.get("is_ml_related")]
        print(f"      -> filtered to {len(papers)}/{before} ML-related papers.")

    if not papers:
        sys.exit("No ML-related papers left after filtering. Try without --ml-only.")

    print("[3/4] Clustering papers into subtopics via Gemini ...")
    subtopics = cluster_into_subtopics(model, args.topic, papers)
    for s in subtopics:
        print(f"      {s['id']}. {s['name']} ({len(s['paper_indices'])} papers)")

    print("[4/4] Writing Excel report ...")
    safe_name = re.sub(r"[^a-z0-9]+", "_", args.topic.lower()).strip("_")
    filename = f"literature_review_{safe_name}.xlsx"
    export_to_excel(args.topic, subtopics, papers, filename)


if __name__ == "__main__":
    main()
