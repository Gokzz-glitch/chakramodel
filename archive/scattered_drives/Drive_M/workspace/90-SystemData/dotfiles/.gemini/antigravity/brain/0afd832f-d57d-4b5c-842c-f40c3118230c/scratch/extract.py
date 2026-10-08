import pptx
import json
import re
import datetime
import os

file_path = r"C:\Users\imgk3\Downloads\ChakraModel_Pitch_Deck (1).pptx"

def extract_pptx(path):
    prs = pptx.Presentation(path)
    slides = []
    all_text = []
    
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text.append(shape.text)
        slides.append({
            "slide_number": i + 1,
            "text": "\n".join(slide_text)
        })
        all_text.extend(slide_text)
    
    combined_text = "\n".join(all_text)
    
    # Extract metrics and values
    pattern = r'\b\d+(?:\.\d+)?(?:%|M|K|B|k|m|b)?\b|\b\$\d+(?:\.\d+)?(?:M|K|B|k|m|b)?\b'
    metrics = list(set(re.findall(pattern, combined_text)))
    
    return {
        "file_name": os.path.basename(path),
        "file_path": path,
        "timestamp_processed": datetime.datetime.now().isoformat(),
        "file_type": "pptx",
        "extracted_content": {
            "slides": slides,
            "all_text": combined_text
        },
        "key_metrics_and_values": metrics
    }

data = extract_pptx(file_path)
out_path = r"m:\chakramodel\ChakraModel_Pitch_Deck (1).pptx_extracted.json"
with open(out_path, "w", encoding="utf-8") as f:
    json.dump(data, f, indent=2, ensure_ascii=False)
print("Saved to", out_path)
