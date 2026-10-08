from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
import os

template_path = r"J:\My Drive\downloads\Project_Review_PPT_Template.pptx"
output_path = r"C:\Users\imgk3\.gemini\antigravity\scratch\digital-academic-doc-workflow\Academic_Document_Workflow_Presentation_Templated.pptx"

prs = Presentation(template_path)

# Let's remove any existing slides in the template to start fresh, 
# or just append to them. Appending is safer, but if the template has dummy slides we might want to delete them.
# For now, let's just append new slides.

def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.placeholders[0].text = title
    slide.placeholders[1].text = subtitle

def add_content_slide(prs, title, content_bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.placeholders[0].text = title
    
    tf = slide.placeholders[1].text_frame
    tf.clear() # clear existing paragraphs if any
    
    for i, bullet in enumerate(content_bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
            
        p.text = bullet['text']
        if 'level' in bullet:
            p.level = bullet['level']

# SLIDE 1: Title Slide
add_title_slide(
    prs, 
    "Digital Academic Document Workflow System", 
    "Project Presentation\n\nPrepared by:\nGOKUL R\nRoll No: 310624148027\nCourse: 2311CSC503R\nEaswari Engineering College"
)

# SLIDE 2: Problem Statement
add_content_slide(prs, "Problem Statement", [
    {'text': 'Manual Process Bottlenecks:', 'level': 0},
    {'text': 'Students must physically visit multiple departments to get approvals.', 'level': 1},
    {'text': 'Lack of Transparency:', 'level': 0},
    {'text': 'Students are unaware of the status of their document requests.', 'level': 1},
    {'text': 'Authenticity & Verification Issues:', 'level': 0},
    {'text': 'Manual signatures can be forged, and physical documents are hard to verify.', 'level': 1},
    {'text': 'Administrative Overhead:', 'level': 0},
    {'text': 'Paper-based workflows consume time, storage space, and administrative effort.', 'level': 1}
])

# SLIDE 3: Proposed Solution
add_content_slide(prs, "Proposed Solution", [
    {'text': 'A completely digitized MERN stack web application to manage the entire lifecycle of academic document requests.', 'level': 0},
    {'text': 'Dynamic Workflow Automation:', 'level': 0},
    {'text': 'Automated routing of requests to Class Advisor, HOD, and Dean based on document type.', 'level': 1},
    {'text': 'Automated Document Generation:', 'level': 0},
    {'text': 'System-generated PDFs with official institutional formatting.', 'level': 1},
    {'text': 'QR Code Verification:', 'level': 0},
    {'text': 'Embedded QR codes link to a public verification portal to ensure authenticity.', 'level': 1}
])

# SLIDE 4: System Architecture
add_content_slide(prs, "System Architecture", [
    {'text': 'Frontend (Client-Side):', 'level': 0},
    {'text': 'React.js with Vite for fast performance and modern UI.', 'level': 1},
    {'text': 'React Router for protected, role-based navigation.', 'level': 1},
    {'text': 'Backend (Server-Side):', 'level': 0},
    {'text': 'Node.js & Express.js for RESTful API development.', 'level': 1},
    {'text': 'Database:', 'level': 0},
    {'text': 'MongoDB with Mongoose ODM for flexible schema design.', 'level': 1},
    {'text': 'Key Libraries:', 'level': 0},
    {'text': 'PDFKit (Document Generation), QRCode (Verification), JWT & bcrypt (Security).', 'level': 1}
])

# SLIDE 5: Approval Workflow
add_content_slide(prs, "Multi-stage Approval Workflow", [
    {'text': 'Configurable state-machine workflow based on Document Type:', 'level': 0},
    {'text': '1. Student Submits Request', 'level': 1},
    {'text': '2. Class Advisor Reviews & Approves', 'level': 1},
    {'text': '3. HOD Approves', 'level': 1},
    {'text': '4. Admin/Dean Finalizes', 'level': 1},
    {'text': '5. System Generates PDF & attaches QR Verification Code', 'level': 1}
])

# SLIDE 6: Key System Features
add_content_slide(prs, "Key System Features", [
    {'text': 'Role-Based Access Control (RBAC):', 'level': 0},
    {'text': 'Strict separation of permissions for Students, Advisors, HODs, and Admins.', 'level': 1},
    {'text': 'Real-time Notifications:', 'level': 0},
    {'text': 'In-app notification center and email alerts at every stage.', 'level': 1},
    {'text': 'Dynamic Document Types:', 'level': 0},
    {'text': 'Admin can create new document types (e.g., Bonafide, NOC) with custom fields.', 'level': 1},
    {'text': 'Public Verification Portal:', 'level': 0},
    {'text': 'Anyone can scan the document QR code to verify its authenticity instantly.', 'level': 1}
])

# SLIDE 7: Conclusion
add_content_slide(prs, "Conclusion & Impact", [
    {'text': 'Efficiency:', 'level': 0},
    {'text': 'Reduces document processing time from days to hours.', 'level': 1},
    {'text': 'Sustainability:', 'level': 0},
    {'text': 'Promotes a paperless, eco-friendly academic environment.', 'level': 1},
    {'text': 'Security:', 'level': 0},
    {'text': 'Eliminates document forgery through cryptographically secure QR verification.', 'level': 1}
])

add_title_slide(prs, "Thank You!", "Any Questions?")

prs.save(output_path)
print(f"Presentation saved successfully to: {output_path}")
