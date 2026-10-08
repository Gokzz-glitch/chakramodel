import sys
import json
import datetime
import re
import os

try:
    from pptx import Presentation
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation

pptx_path = r"C:\Users\imgk3\Downloads\ChakraModel_Final_Pitch.pptx"
json_path = r"m:\chakramodel\ChakraModel_Final_Pitch.pptx_extracted.json"

extracted_content = {"slides": []}
key_metrics = []

prs = Presentation(pptx_path)
for i, slide in enumerate(prs.slides):
    slide_content = {"slide_number": i + 1, "text": []}
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            text = shape.text.strip()
            if text:
                slide_content["text"].append(text)
                # find numbers/percentages
                metrics = re.findall(r'\b\d+(?:\.\d+)?%?|\$\d+(?:\.\d+)?[MBK]?\b', text)
                key_metrics.extend(metrics)
    extracted_content["slides"].append(slide_content)

output = {
    "file_name": "ChakraModel_Final_Pitch.pptx",
    "file_path": pptx_path,
    "timestamp_processed": datetime.datetime.now().isoformat(),
    "file_type": "pptx",
    "extracted_content": extracted_content,
    "key_metrics_and_values": list(set(key_metrics))
}

os.makedirs(os.path.dirname(json_path), exist_ok=True)
with open(json_path, "w", encoding="utf-8") as f:
    json.dump(output, f, indent=4)

print(f"Saved {json_path}")
